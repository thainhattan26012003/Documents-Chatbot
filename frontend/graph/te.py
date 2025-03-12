import re
import json
from docx import Document

# Định nghĩa ánh xạ footnote (dựa theo file mẫu)
footnote_mapping = {
    "1": "ten_co_quan_chu_quan",
    "2": "ten_co_quan_cap_giay",
    "3": "chu_viet_tat",
    "4": "dia_danh",
    "5": "ho_va_chuc_vu",
    "6": "noi_nghi_phep",
    "7": "thoi_gian_nghi",
    "8": "nguoi_cap_giay",
    "9": "chu_viet_tat_don_vi"
}

def extract_placeholders_with_keys(docx_file):
    # Mở file DOCX
    doc = Document(docx_file)
    placeholders = {}
    placeholder_counter = 1  # dùng để đánh số các placeholder nếu không có footnote số

    # Pattern: tìm chuỗi dấu chấm (ít nhất 3 dấu) và tùy chọn theo sau bởi số
    pattern = re.compile(r'(\.{3,})(\d+)?')
    
    for para in doc.paragraphs:
        for match in pattern.finditer(para.text):
            dots = match.group(1)
            digit = match.group(2)
            if digit:
                key = footnote_mapping.get(digit, f"placeholder_{placeholder_counter}")
            else:
                key = f"placeholder_{placeholder_counter}"
            placeholders[key] = dots  # giá trị ban đầu có thể là chuỗi placeholder, bạn có thể để rỗng nếu muốn
            placeholder_counter += 1

    return placeholders

if __name__ == "__main__":
    file_path = "GIAY NGHI PHEP.docx"  # đường dẫn file DOCX
    ph_dict = extract_placeholders_with_keys(file_path)
    # Xuất kết quả thành file JSON
    json_output = json.dumps(ph_dict, ensure_ascii=False, indent=4)
    print("JSON output:")
    print(json_output)
    # Nếu cần ghi ra file:
    with open("placeholders.json", "w", encoding="utf-8") as f:
        f.write(json_output)
